"""AC04/AC05/T58: the six workbook reference employees, component by component.

Inputs and expected displayed values are read from the original workbook (openpyxl). The only
hand-written fixture is the mapping of each employee's free-text current vehicle / alternative
to catalogue rows, because the workbook names vehicles by commercial name only.
Precise raw values are cross-checked against Agent Spec §8.
"""

from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal
from typing import Any

import openpyxl
import pytest

from app.business.comparisons.builder import build_comparison
from app.business.display_rounding import display_delta, display_round
from app.domain.arrangements import (
    AlternativeSpec,
    AlternativeVehicleSpec,
    BaselineSpec,
    CurrentAllowanceSpec,
    CurrentCompanyCarSpec,
    WaiverSpec,
)
from app.domain.comparisons import ComparisonOutcome, ScenarioStatus
from app.domain.employee_profiles import SalaryFacts
from app.domain.provenance import SourceType
from app.domain.rules import EntitlementGroup
from app.domain.vehicles import Propulsion, Track
from tests.paths import REPO_ROOT
from tests.unit.business.helpers import rules, usable

WORKBOOK = REPO_ROOT / "data" / "source" / "challenge-data.xlsx"
PROPULSION = {
    "בנזין": Propulsion.PETROL,
    "היברידי": Propulsion.HYBRID,
    "פלאג-אין": Propulsion.PLUG_IN_HYBRID,
    "חשמלי": Propulsion.ELECTRIC,
}

# Free-text workbook descriptions -> catalogue rows / arrangement kinds (explicit test fixture).
CURRENT = {
    "E-001": "CAR-09",  # COROLLA CROSS (2025)
    "E-002": "CAR-05",  # NIRO (2025)
    "E-003": "allowance",  # ויתור - אחזקת רכב 1500 ₪
    "E-004": "CAR-12",  # TUCSON (2025)
    "E-005": "CAR-13",  # MAZDA 3 (2025)
    "E-006": "CAR-15",  # OCTAVIA (2025)
}
ALTERNATIVE = {
    "E-001": "CAR-10",  # ATTO 3
    "E-002": "CAR-19",  # CARNIVAL, family track
    "E-003": "CAR-06",  # MG4
    "E-004": "CAR-11",  # EV3
    "E-005": "waiver",  # ויתור - אחזקת רכב 2000 ₪
    "E-006": "CAR-14",  # MODEL 3
}

# Agent Spec §8: exact current net, exact alternative net, exact difference (3 decimals).
SPEC_RAW = {
    "E-001": ("11031.819", "11692.817", "660.998"),
    "E-002": ("8599.733", "5077.844", "-3521.889"),
    "E-003": ("15047.338", "13108.657", "-1938.681"),
    "E-004": ("10692.845", "11899.807", "1206.962"),
    "E-005": ("14244.019", "17806.288", "3562.269"),
    "E-006": ("17072.533", "17794.233", "721.700"),
}

# Documented workbook exceptions (SRC-07): workbook delta cell vs approved precise rounding.
DELTA_EXCEPTIONS = {"E-003": (-1938, -1939), "E-006": (721, 722)}


def _dec(cell: Any) -> Decimal:
    if isinstance(cell, bool) or cell is None:
        raise AssertionError(f"unexpected workbook cell {cell!r}")
    return Decimal(str(cell))


def _rows(sheet: Any) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(h).strip() for h in rows[0]]
    return [dict(zip(header, row, strict=False)) for row in rows[1:] if row[0] is not None]


@dataclass(frozen=True)
class Workbook:
    employees: dict[str, dict[str, Any]]
    calculations: dict[str, dict[str, Any]]
    catalogue: dict[str, dict[str, Any]]
    occupations: dict[str, dict[str, Any]]


@pytest.fixture(scope="module")
def workbook() -> Workbook:
    book = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        return Workbook(
            employees={r["מזהה"]: r for r in _rows(book["עובדים לדוגמה"])},
            calculations={r["מזהה"]: r for r in _rows(book["חישובים"])},
            catalogue={r["מזהה"]: r for r in _rows(book["קטלוג רכבים"])},
            occupations={r["עיסוק"]: r for r in _rows(book["עיסוקים"])},
        )
    finally:
        book.close()


def _salary(employee: dict[str, Any]) -> SalaryFacts:
    hundred = Decimal(100)
    return SalaryFacts(
        base_gross=usable(_dec(employee["שכר בסיס"])),
        commissions_gross=usable(_dec(employee["עמלות"])),
        bonus_gross=usable(_dec(employee["בונוס"])),
        pension_employee_rate=usable(_dec(employee["פנסיה עובד %"]) / hundred),
        study_fund_employee_rate=usable(_dec(employee["קרן השתלמות עובד %"]) / hundred),
        credit_points=usable(_dec(employee["נקודות זיכוי"])),
    )


def _baseline(employee_id: str, wb: Workbook) -> BaselineSpec:
    employee = wb.employees[employee_id]
    current = CURRENT[employee_id]
    if current == "allowance":
        allowance = _dec(wb.occupations[employee["עיסוק"]]["אחזקת רכב בוויתור (₪)"])
        return CurrentAllowanceSpec(allowance=usable(allowance))
    car = wb.catalogue[current]
    return CurrentCompanyCarSpec(
        track=Track.REGULAR,
        propulsion=usable(PROPULSION[car["הנעה"]], SourceType.CATALOGUE),
        value_of_use=usable(_dec(car["שווי שימוש חודשי"]), SourceType.TAX_DATASET),
        participation=usable(_dec(employee["השתתפות היום"])),
        catalogue_id=current,
    )


def _alternative(employee_id: str, wb: Workbook) -> AlternativeSpec:
    employee = wb.employees[employee_id]
    target = ALTERNATIVE[employee_id]
    if target == "waiver":
        allowance = _dec(wb.occupations[employee["עיסוק"]]["אחזקת רכב בוויתור (₪)"])
        return WaiverSpec(allowance=usable(allowance, SourceType.CATALOGUE))
    car = wb.catalogue[target]
    family = car["קבוצה"] == "משפחתי"
    group = EntitlementGroup(employee["קבוצה"])
    return AlternativeVehicleSpec(
        option_key=f"vehicle:{target}",
        catalogue_id=target,
        track=Track.FAMILY if family else Track.REGULAR,
        propulsion=PROPULSION[car["הנעה"]],
        value_of_use=usable(_dec(car["שווי שימוש חודשי"]), SourceType.TAX_DATASET),
        participation=usable(_dec(car["השתתפות עובד ברירת מחדל (דמה)"]), SourceType.CATALOGUE),
        monthly_company_cost=_dec(car["עלות חודשית לחברה (דמה)"]),
        group_budget=usable(rules().budgets[group], SourceType.RULES),
    )


def _compare(employee_id: str, wb: Workbook) -> ComparisonOutcome:
    return build_comparison(
        _salary(wb.employees[employee_id]),
        _baseline(employee_id, wb),
        [_alternative(employee_id, wb)],
        rules(),
    )


def _three(amount: Decimal) -> str:
    return str(amount.quantize(Decimal("0.001"), rounding=ROUND_HALF_UP))


EMPLOYEES = sorted(CURRENT)


def test_workbook_has_exactly_the_six_reference_employees(workbook: Workbook) -> None:
    assert sorted(workbook.employees) == EMPLOYEES
    assert sorted(workbook.calculations) == EMPLOYEES


@pytest.mark.parametrize("employee_id", EMPLOYEES)
def test_reference_components_match_workbook(employee_id: str, workbook: Workbook) -> None:
    expected = workbook.calculations[employee_id]
    outcome = _compare(employee_id, workbook)
    assert outcome.baseline is not None and outcome.baseline.result is not None
    alternative = outcome.alternatives[0]
    assert alternative.status is ScenarioStatus.CALCULATED and alternative.result is not None
    current, alt = outcome.baseline.result, alternative.result

    assert display_round(current.net) == _dec(expected["נטו היום"])
    assert display_round(alt.net) == _dec(expected["נטו בחלופה"])
    assert display_round(current.income_tax) == _dec(expected["מס היום"])
    assert display_round(alt.income_tax) == _dec(expected["מס בחלופה"])
    assert display_round(current.combined_insurance) == _dec(expected["ביטוח לאומי+בריאות היום"])
    assert display_round(alt.combined_insurance) == _dec(expected["ביטוח לאומי+בריאות בחלופה"])
    assert display_round(alt.electric_gross_up) == _dec(expected["החזר חשמל ברוטו"])
    assert display_round(alt.family_gap_contribution) == _dec(expected["חלק העובד במסלול המשפחתי"])


@pytest.mark.parametrize("employee_id", EMPLOYEES)
def test_reference_raw_values_match_agent_spec(employee_id: str, workbook: Workbook) -> None:
    outcome = _compare(employee_id, workbook)
    assert outcome.baseline is not None and outcome.baseline.result is not None
    alternative = outcome.alternatives[0]
    assert alternative.result is not None and alternative.delta is not None
    current_net, alternative_net, delta = SPEC_RAW[employee_id]
    assert _three(outcome.baseline.result.net) == current_net
    assert _three(alternative.result.net) == alternative_net
    assert _three(alternative.delta) == delta


@pytest.mark.parametrize("employee_id", EMPLOYEES)
def test_delta_display_rounds_the_precise_difference(employee_id: str, workbook: Workbook) -> None:
    outcome = _compare(employee_id, workbook)
    assert outcome.baseline is not None and outcome.baseline.result is not None
    alternative = outcome.alternatives[0]
    assert alternative.result is not None and alternative.delta is not None
    workbook_delta = _dec(workbook.calculations[employee_id]["הפרש"])
    shown = display_round(alternative.delta)
    assert shown == display_delta(alternative.result.net, outcome.baseline.result.net)
    if employee_id in DELTA_EXCEPTIONS:
        workbook_value, approved_value = DELTA_EXCEPTIONS[employee_id]
        assert workbook_delta == workbook_value  # documented source exception still present
        assert shown == approved_value
        assert shown != workbook_delta
    else:
        assert shown == workbook_delta


def test_documented_exceptions_come_from_rounded_net_subtraction(workbook: Workbook) -> None:
    """SRC-07: the two workbook cells equal (rounded alt net - rounded current net)."""
    for employee_id in DELTA_EXCEPTIONS:
        row = workbook.calculations[employee_id]
        assert _dec(row["הפרש"]) == _dec(row["נטו בחלופה"]) - _dec(row["נטו היום"])


def test_allowance_to_car_transition_leaves_no_residual_allowance(workbook: Workbook) -> None:
    """AC13 (E-003): the old allowance is removed in the vehicle alternative."""
    outcome = _compare("E-003", workbook)
    assert outcome.baseline is not None and outcome.baseline.result is not None
    alternative = outcome.alternatives[0].result
    assert alternative is not None
    assert outcome.baseline.result.allowance == Decimal(1500)
    assert alternative.allowance == 0


def test_waiver_alternative_has_only_the_occupation_allowance(workbook: Workbook) -> None:
    """E-005 / T22: waiver -> A from the occupation table; U, C, F, E zero; A not pensionable."""
    outcome = _compare("E-005", workbook)
    assert outcome.baseline is not None and outcome.baseline.result is not None
    waiver = outcome.alternatives[0].result
    assert waiver is not None
    assert waiver.allowance == Decimal(2000)
    assert (waiver.value_of_use, waiver.ordinary_participation) == (0, 0)
    assert (waiver.family_gap_contribution, waiver.electric_gross_up) == (0, 0)
    assert waiver.pension_employee == outcome.baseline.result.pension_employee
