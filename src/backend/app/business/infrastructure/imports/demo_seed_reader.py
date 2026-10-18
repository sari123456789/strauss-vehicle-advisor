"""`DemoSeedSourceReader`: workbook example-employee sheet + `data/demo/demo-accounts.json`.

The workbook is verified against `data/source-manifest.json` before parsing (openpyxl, no text
encoding). Columns are located by exact Hebrew header. The fictional ID-number column is read
past and never returned. Money and child counts must be integer cells. Percentages and credit
points may be non-integer cells, which openpyxl returns as binary floats; each such value is
converted with `Decimal(repr(value))` only if it round-trips exactly and has at most 6
decimal places, otherwise the seed fails (no silent rounding). Rates are then divided by 100
in Decimal. Occupation names map to IDs through the derived policy (decision R6).
"""

import hashlib
import io
import json
from collections.abc import Mapping
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID

import openpyxl

from app.business.application.ports.demo_seed import (
    DemoCurrentArrangement,
    DemoEmployeeSeed,
    DemoSeedSource,
    DemoStaffSeed,
)
from app.business.application.ports.reference_data import SourceFingerprint
from app.business.domain.arrangements import ArrangementKind
from app.business.domain.rules import EntitlementGroup
from app.business.infrastructure.imports.source_document import SourceManifest
from app.contracts.tool_context import Role

WORKBOOK_KIND = "challenge_workbook_xlsx"
SEED_KIND = "demo_seed_mapping_json"
EMPLOYEES_SHEET = "עובדים לדוגמה"
EMPLOYEES_HEADER: tuple[str, ...] = (
    "מזהה",
    "שם",
    "ת״ז (פיקטיבית)",
    "עיסוק",
    "קבוצה",
    "שכר בסיס",
    "עמלות",
    "בונוס",
    "פנסיה עובד %",
    "קרן השתלמות עובד %",
    "נקודות זיכוי",
    "ילדים עד 23",
    "רכב היום",
    "השתתפות היום",
    "השאלה",
)
HUNDRED = Decimal(100)
MAX_FRACTION_DIGITS = 6


class DemoSeedFormatError(Exception):
    """The seed mapping or workbook does not satisfy the seed contract. Nothing is seeded."""

    code = "SOURCE_FORMAT"


def _exact_decimal(value: object, where: str) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise DemoSeedFormatError(f"{where}: expected a number, got {value!r}")
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        result = Decimal(repr(value))
        exponent = result.as_tuple().exponent
        if (
            not result.is_finite()
            or float(result) != value
            or not isinstance(exponent, int)
            or -exponent > MAX_FRACTION_DIGITS
        ):
            raise DemoSeedFormatError(f"{where}: {value!r} is not an exact short decimal")
        return result
    raise DemoSeedFormatError(f"{where}: expected a number, got {value!r}")


def _integer(value: object, where: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int):
        raise DemoSeedFormatError(f"{where}: expected an integer cell, got {value!r}")
    return value


def _text(value: object, where: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise DemoSeedFormatError(f"{where}: expected a text cell, got {value!r}")
    return value.strip()


def _require(mapping: Mapping[str, Any], keys: set[str], where: str) -> None:
    found = set(mapping)
    if not keys <= found:
        raise DemoSeedFormatError(f"{where}: missing keys {sorted(keys - found)}")


class FileDemoSeedSourceReader:
    def __init__(
        self, repo_root: Path, manifest_path: Path, policy_path: Path, seed_path: Path
    ) -> None:
        self._repo_root = repo_root
        self._manifest_path = manifest_path
        self._policy_path = policy_path
        self._seed_path = seed_path

    def read(self) -> DemoSeedSource:
        raw_seed = self._seed_path.read_bytes()
        seed = self._json(raw_seed, self._seed_path.name)
        policy = self._json(self._policy_path.read_bytes(), self._policy_path.name)
        ids_by_name = policy.get("occupationIds", {}).get("byName")
        if not isinstance(ids_by_name, Mapping):
            raise DemoSeedFormatError("derived policy is missing occupationIds.byName")

        _require(seed, {"seedVersion", "uuidNamespace", "employees", "staff"}, "seed")
        try:
            namespace = UUID(seed["uuidNamespace"])
        except (ValueError, TypeError, AttributeError):
            raise DemoSeedFormatError("seed: uuidNamespace must be a UUID") from None
        mapping = {e["employeeRef"]: e for e in seed["employees"]}
        if len(mapping) != len(seed["employees"]):
            raise DemoSeedFormatError("seed: duplicate employeeRef")

        document = SourceManifest.load(self._repo_root, self._manifest_path).read(WORKBOOK_KIND)
        rows = self._employee_rows(document.raw)
        if set(rows) != set(mapping):
            raise DemoSeedFormatError(
                f"seed employees {sorted(mapping)} differ from workbook employees {sorted(rows)}"
            )
        employees = tuple(
            self._employee(ref, rows[ref], mapping[ref], ids_by_name) for ref in sorted(rows)
        )
        staff = tuple(self._staff(s) for s in seed["staff"])
        return DemoSeedSource(
            seed_version=str(seed["seedVersion"]),
            uuid_namespace=namespace,
            seed_file=SourceFingerprint(
                kind=SEED_KIND,
                path=self._display_path(self._seed_path),
                sha256=hashlib.sha256(raw_seed).hexdigest(),
                size_bytes=len(raw_seed),
            ),
            workbook=document.fingerprint,
            employees=employees,
            staff=staff,
        )

    def _display_path(self, path: Path) -> str:
        try:
            return path.relative_to(self._repo_root).as_posix()
        except ValueError:
            return path.name

    @staticmethod
    def _json(raw: bytes, name: str) -> dict[str, Any]:
        try:
            content = json.loads(raw.decode("utf-8"), parse_float=Decimal)
        except (UnicodeDecodeError, json.JSONDecodeError) as error:
            raise DemoSeedFormatError(f"{name}: not valid UTF-8 JSON ({error})") from error
        if not isinstance(content, dict):
            raise DemoSeedFormatError(f"{name}: top-level JSON value must be an object")
        return content

    @staticmethod
    def _employee_rows(raw: bytes) -> dict[str, dict[str, Any]]:
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as error:  # openpyxl raises several unrelated types for corrupt files
            raise DemoSeedFormatError(f"workbook cannot be opened ({error})") from error
        try:
            if EMPLOYEES_SHEET not in workbook.sheetnames:
                raise DemoSeedFormatError(f"workbook has no sheet named {EMPLOYEES_SHEET!r}")
            cells = list(workbook[EMPLOYEES_SHEET].iter_rows(values_only=True))
        finally:
            workbook.close()
        header = tuple(h.strip() if isinstance(h, str) else h for h in cells[0])
        while header and header[-1] is None:
            header = header[:-1]
        if header != EMPLOYEES_HEADER:
            raise DemoSeedFormatError(f"{EMPLOYEES_SHEET}: unexpected header {list(header)}")
        rows: dict[str, dict[str, Any]] = {}
        for number, values in enumerate(cells[1:], start=2):
            if all(v is None for v in values):
                continue
            row = dict(zip(EMPLOYEES_HEADER, values, strict=False))
            ref = _text(row["מזהה"], f"{EMPLOYEES_SHEET} row {number}")
            if ref in rows:
                raise DemoSeedFormatError(f"{EMPLOYEES_SHEET}: duplicate employee {ref}")
            rows[ref] = row
        return rows

    @staticmethod
    def _employee(
        ref: str, row: dict[str, Any], mapping: Mapping[str, Any], ids_by_name: Mapping[str, Any]
    ) -> DemoEmployeeSeed:
        where = f"{EMPLOYEES_SHEET} {ref}"
        _require(mapping, {"login", "currentArrangement"}, f"seed {ref}")
        occupation_name = _text(row["עיסוק"], f"{where} עיסוק")
        occupation_id = ids_by_name.get(occupation_name)
        if not isinstance(occupation_id, str):
            raise DemoSeedFormatError(f"{where}: unmapped occupation {occupation_name!r}")
        try:
            group = EntitlementGroup(_text(row["קבוצה"], f"{where} קבוצה"))
            kind = ArrangementKind(mapping["currentArrangement"]["kind"])
        except (ValueError, KeyError, TypeError) as error:
            raise DemoSeedFormatError(f"{where}: {error}") from None
        current_map = mapping["currentArrangement"]
        workbook_text = _text(row["רכב היום"], f"{where} רכב היום")
        if current_map.get("workbookText") != workbook_text:
            raise DemoSeedFormatError(
                f"{where}: workbook current car {workbook_text!r} differs from the seed mapping "
                f"{current_map.get('workbookText')!r}"
            )
        participation = _integer(row["השתתפות היום"], f"{where} השתתפות היום")
        if kind is ArrangementKind.COMPANY_CAR:
            current = DemoCurrentArrangement(
                kind=kind,
                catalogue_id=current_map.get("catalogueId"),
                expected_commercial_name=workbook_text,
                participation=Decimal(participation),
            )
        elif kind is ArrangementKind.ALLOWANCE:
            if participation != 0:
                raise DemoSeedFormatError(f"{where}: allowance with non-zero participation")
            allowance = current_map.get("allowance")
            if not isinstance(allowance, str) or not allowance.isdigit():
                raise DemoSeedFormatError(f"{where}: allowance must be a whole-shekel string")
            current = DemoCurrentArrangement(kind=kind, allowance=Decimal(allowance))
        else:
            current = DemoCurrentArrangement(kind=kind)
        return DemoEmployeeSeed(
            employee_ref=ref,
            login=str(mapping["login"]),
            display_name=_text(row["שם"], f"{where} שם"),
            occupation_id=occupation_id,
            entitlement_group=group,
            base_gross=Decimal(_integer(row["שכר בסיס"], f"{where} שכר בסיס")),
            commissions_gross=Decimal(_integer(row["עמלות"], f"{where} עמלות")),
            bonus_gross=Decimal(_integer(row["בונוס"], f"{where} בונוס")),
            pension_employee_rate=_exact_decimal(row["פנסיה עובד %"], f"{where} פנסיה") / HUNDRED,
            study_fund_employee_rate=(
                _exact_decimal(row["קרן השתלמות עובד %"], f"{where} קרן השתלמות") / HUNDRED
            ),
            credit_points=_exact_decimal(row["נקודות זיכוי"], f"{where} נקודות זיכוי"),
            qualifying_children_count=_integer(row["ילדים עד 23"], f"{where} ילדים"),
            current=current,
        )

    @staticmethod
    def _staff(item: Mapping[str, Any]) -> DemoStaffSeed:
        _require(item, {"login", "roles", "displayName"}, "seed staff")
        try:
            roles = frozenset(Role(r) for r in item["roles"])
        except (ValueError, TypeError) as error:
            raise DemoSeedFormatError(f"seed staff {item.get('login')!r}: {error}") from None
        return DemoStaffSeed(
            login=str(item["login"]), roles=roles, display_name=str(item["displayName"])
        )
