"""Challenge workbook: vehicle catalogue and occupation table (openpyxl, no text encoding).

Sheets and columns are located by their exact Hebrew names, never by position. Identifier codes
must be text cells (a numeric cell would already have lost its leading zeros); money and years
must be integer cells (no float -> Decimal conversion). The catalogue omits the vehicle-type
code, so each vehicle's complete tax key comes from the explicit per-vehicle mapping in the
derived policy (SRC-03); an unmapped vehicle fails the import. Occupation IDs come from the
derived policy mapping (decision R6). Example-employee and calculation sheets are not imported.
"""

import io
import re
from collections.abc import Mapping
from decimal import Decimal
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from app.business.application.ports.reference_data import (
    CatalogueSource,
    ImportReport,
    JsonScalar,
    SourceRow,
    TracedOccupation,
    TracedVehicle,
)
from app.business.domain.errors import DomainError
from app.business.domain.occupations import OccupationPolicy
from app.business.domain.rules import EntitlementGroup
from app.business.domain.vehicles import CatalogueVehicle, Propulsion, Track, VehicleTaxKey
from app.business.infrastructure.imports.source_document import SourceDocument

CATALOGUE_SHEET = "קטלוג רכבים"
OCCUPATIONS_SHEET = "עיסוקים"
DESCRIPTION_SHEET = "כללים ונתונים"

CATALOGUE_HEADER: tuple[str, ...] = (
    "מזהה",
    "קבוצה",
    "תוצר",
    "קוד תוצר",
    "דגם (כינוי מסחרי)",
    "קוד דגם",
    "תאור דגם",
    "שנת רישום",
    "הנעה",
    "מחיר מחירון",
    "שווי שימוש חודשי",
    "עלות חודשית לחברה (דמה)",
    "השתתפות עובד ברירת מחדל (דמה)",
)
OCCUPATIONS_HEADER: tuple[str, ...] = (
    "עיסוק",
    "קבוצת זכאות",
    "רכב חובה",
    "מותר לוותר",
    "אחזקת רכב בוויתור (₪)",
)
FAMILY_GROUP_LABEL = "משפחתי"
YES, NO = "כן", "לא"

PROPULSION_BY_LABEL: Mapping[str, Propulsion] = {
    "בנזין": Propulsion.PETROL,
    "היברידי": Propulsion.HYBRID,
    "פלאג-אין": Propulsion.PLUG_IN_HYBRID,
    "חשמלי": Propulsion.ELECTRIC,
}
"""Translation of the rules JSON `driveField` labels (also used by the catalogue) to categories."""

_CODE5 = re.compile(r"^\d{5}$")
_CLAIMED_COUNTS = re.compile(r"(\d+)\s*רכבים.*?\+\s*(\d+)")


class WorkbookFormatError(Exception):
    """The workbook does not satisfy the import contract. Nothing is imported."""

    code = "SOURCE_FORMAT"


class _Row:
    def __init__(self, sheet: str, number: int, header: tuple[str, ...], cells: tuple[Any, ...]):
        self.sheet, self.number = sheet, number
        self._cells = dict(zip(header, cells, strict=True))

    def _where(self, column: str) -> str:
        return f"{self.sheet} row {self.number}, {column}"

    def text(self, column: str) -> str:
        value = self._cells[column]
        if not isinstance(value, str) or not value.strip():
            raise WorkbookFormatError(f"{self._where(column)}: expected a text cell, got {value!r}")
        return value.strip()

    def code5(self, column: str) -> str:
        value = self._cells[column]
        if not isinstance(value, str):
            raise WorkbookFormatError(
                f"{self._where(column)}: code must be a text cell (leading zeros), got {value!r}"
            )
        if not _CODE5.fullmatch(value.strip()):
            raise WorkbookFormatError(f"{self._where(column)}: invalid 5-digit code {value!r}")
        return value.strip()

    def integer(self, column: str) -> int:
        value = self._cells[column]
        if isinstance(value, bool) or not isinstance(value, int):
            raise WorkbookFormatError(
                f"{self._where(column)}: expected an integer cell, got {value!r}"
            )
        return value

    def money(self, column: str) -> Decimal:
        return Decimal(self.integer(column))

    def yes_no(self, column: str) -> bool:
        value = self.text(column)
        if value not in (YES, NO):
            raise WorkbookFormatError(f"{self._where(column)}: expected {YES}/{NO}, got {value!r}")
        return value == YES


def _rows(sheet: Worksheet, expected: tuple[str, ...]) -> list[tuple[_Row, SourceRow]]:
    rows = sheet.iter_rows(values_only=True)
    try:
        header = tuple(rows.__next__())
    except StopIteration:
        raise WorkbookFormatError(f"sheet {sheet.title!r} is empty") from None
    trimmed = tuple(h.strip() if isinstance(h, str) else h for h in header)
    while trimmed and trimmed[-1] is None:
        trimmed = trimmed[:-1]
    if trimmed != expected:
        raise WorkbookFormatError(
            f"sheet {sheet.title!r}: unexpected header {list(trimmed)}; expected {list(expected)}"
        )
    result: list[tuple[_Row, SourceRow]] = []
    for number, cells in enumerate(rows, start=2):
        values = tuple(cells[: len(expected)]) + (None,) * (len(expected) - len(cells))
        extra = cells[len(expected) :]
        if all(v is None for v in values) and all(v is None for v in extra):
            continue
        if any(v is not None for v in extra):
            raise WorkbookFormatError(f"sheet {sheet.title!r} row {number}: unexpected extra cells")
        for value in values:
            if value is not None and not isinstance(value, str | int):
                raise WorkbookFormatError(
                    f"sheet {sheet.title!r} row {number}: unsupported cell value {value!r}"
                )
        source_values: tuple[JsonScalar, ...] = values
        result.append(
            (_Row(sheet.title, number, expected, values), SourceRow(number, source_values))
        )
    return result


def _sheet(workbook: Any, name: str) -> Worksheet:
    if name not in workbook.sheetnames:
        raise WorkbookFormatError(f"workbook has no sheet named {name!r}")
    sheet: Worksheet = workbook[name]
    return sheet


def _group(label: str, where: str) -> tuple[Track, EntitlementGroup | None]:
    if label == FAMILY_GROUP_LABEL:
        return Track.FAMILY, None
    try:
        return Track.REGULAR, EntitlementGroup(label)
    except ValueError:
        raise WorkbookFormatError(f"{where}: unknown group {label!r}") from None


def _vehicles(
    sheet: Worksheet, tax_year: str, type_codes: Mapping[str, str]
) -> list[TracedVehicle]:
    traced: list[TracedVehicle] = []
    for row, source in _rows(sheet, CATALOGUE_HEADER):
        catalogue_id = row.text("מזהה")
        where = f"{CATALOGUE_SHEET} row {row.number} ({catalogue_id})"
        track, group = _group(row.text("קבוצה"), where)
        type_code = type_codes.get(catalogue_id)
        if type_code is None:
            raise WorkbookFormatError(
                f"{where}: no explicit vehicle-type mapping in the derived policy (SRC-03)"
            )
        propulsion_label = row.text("הנעה")
        propulsion = PROPULSION_BY_LABEL.get(propulsion_label)
        if propulsion is None:
            raise WorkbookFormatError(f"{where}: unknown propulsion {propulsion_label!r}")
        try:
            vehicle = CatalogueVehicle(
                catalogue_id=catalogue_id,
                track=track,
                group=group,
                manufacturer_name=row.text("תוצר"),
                commercial_name=row.text("דגם (כינוי מסחרי)"),
                model_description=row.text("תאור דגם"),
                tax_key=VehicleTaxKey(
                    tax_year=tax_year,
                    registration_year=str(row.integer("שנת רישום")),
                    vehicle_type_code=type_code,
                    manufacturer_code=row.code5("קוד תוצר"),
                    model_code=row.code5("קוד דגם"),
                ),
                propulsion=propulsion,
                list_price=row.money("מחיר מחירון"),
                catalogue_value_of_use=row.money("שווי שימוש חודשי"),
                monthly_company_cost=row.money("עלות חודשית לחברה (דמה)"),
                default_participation=row.money("השתתפות עובד ברירת מחדל (דמה)"),
            )
        except DomainError as error:
            raise WorkbookFormatError(f"{where}: {error}") from error
        traced.append(TracedVehicle(vehicle=vehicle, source=source))
    unused = sorted(set(type_codes) - {t.vehicle.catalogue_id for t in traced})
    if unused:
        raise WorkbookFormatError(f"vehicle-type mapping lists unknown catalogue IDs: {unused}")
    return traced


def _occupations(sheet: Worksheet, ids_by_name: Mapping[str, str]) -> list[TracedOccupation]:
    traced: list[TracedOccupation] = []
    for row, source in _rows(sheet, OCCUPATIONS_HEADER):
        name = row.text("עיסוק")
        where = f"{OCCUPATIONS_SHEET} row {row.number} ({name})"
        occupation_id = ids_by_name.get(name)
        if occupation_id is None:
            raise WorkbookFormatError(f"{where}: occupation has no ID in the derived policy (R6)")
        track, group = _group(row.text("קבוצת זכאות"), where)
        if group is None:
            raise WorkbookFormatError(f"{where}: an occupation must have a regular group")
        try:
            occupation = OccupationPolicy(
                occupation_id=occupation_id,
                name_he=name,
                group=group,
                mandatory_car=row.yes_no("רכב חובה"),
                waiver_allowed=row.yes_no("מותר לוותר"),
                waiver_allowance=row.money("אחזקת רכב בוויתור (₪)"),
            )
        except DomainError as error:
            raise WorkbookFormatError(f"{where}: {error}") from error
        traced.append(TracedOccupation(occupation=occupation, source=source))
    missing = sorted(set(ids_by_name) - {t.occupation.name_he for t in traced})
    if missing:
        raise WorkbookFormatError(f"mapped occupations missing from the workbook: {missing}")
    return traced


def _count_warnings(workbook: Any, regular: int, family: int) -> list[str]:
    """SRC-02: the description text claims different counts than the actual rows."""
    if DESCRIPTION_SHEET not in workbook.sheetnames:
        return []
    for cells in workbook[DESCRIPTION_SHEET].iter_rows(values_only=True):
        if len(cells) >= 2 and cells[0] == "קטלוג הרכבים" and isinstance(cells[1], str):
            match = _CLAIMED_COUNTS.search(cells[1])
            if match and (int(match.group(1)), int(match.group(2))) != (regular, family):
                return [
                    f"SRC-02: workbook description states {match.group(1)} + {match.group(2)} "
                    f"vehicles; actual rows are {regular} regular + {family} family"
                ]
    return []


def _policy_mapping(policy: Mapping[str, Any], path: tuple[str, ...]) -> Any:
    current: Any = policy
    for key in path:
        if not isinstance(current, Mapping) or key not in current:
            raise WorkbookFormatError(f"derived policy is missing {'.'.join(path)}")
        current = current[key]
    return current


def read_catalogue(
    document: SourceDocument, policy: Mapping[str, Any], policy_sha256: str
) -> CatalogueSource:
    tax_year = _policy_mapping(policy, ("catalogueTaxKeys", "taxYear"))
    type_codes = _policy_mapping(policy, ("catalogueTaxKeys", "vehicleTypeCodeByCatalogueId"))
    ids_by_name = _policy_mapping(policy, ("occupationIds", "byName"))
    for name, mapping in (("vehicle-type", type_codes), ("occupation-ID", ids_by_name)):
        if not isinstance(mapping, Mapping) or not all(
            isinstance(k, str) and isinstance(v, str) for k, v in mapping.items()
        ):
            raise WorkbookFormatError(f"derived policy {name} mapping must map strings to strings")
    if not isinstance(tax_year, str):
        raise WorkbookFormatError("derived policy catalogueTaxKeys.taxYear must be a string")
    if len(set(ids_by_name.values())) != len(ids_by_name):
        raise WorkbookFormatError("derived policy reuses an occupation ID")

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(document.raw), read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises several unrelated types for corrupt files
        raise WorkbookFormatError(f"{document.fingerprint.path}: cannot open ({error})") from error
    try:
        vehicles = _vehicles(_sheet(workbook, CATALOGUE_SHEET), tax_year, type_codes)
        occupations = _occupations(_sheet(workbook, OCCUPATIONS_SHEET), ids_by_name)
        family = sum(1 for t in vehicles if t.vehicle.track is Track.FAMILY)
        warnings = _count_warnings(workbook, len(vehicles) - family, family)
    finally:
        workbook.close()

    fingerprint = document.fingerprint
    return CatalogueSource(
        workbook=fingerprint,
        policy_sha256=policy_sha256,
        vehicles=tuple(vehicles),
        occupations=tuple(occupations),
        report=ImportReport(
            kind=fingerprint.kind,
            sha256=fingerprint.sha256,
            record_count=len(vehicles) + len(occupations),
            warnings=tuple(warnings),
            details={
                "regularVehicles": len(vehicles) - family,
                "familyVehicles": family,
                "occupations": len(occupations),
            },
        ),
    )
